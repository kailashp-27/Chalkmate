import io
import os
import time
from uuid import UUID, uuid4

import psycopg
import pytest
from conftest import auth
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from chalkmate import assessment_models as m
from chalkmate.assessment_rules import validate_marks, validate_suggestion
from chalkmate.assessment_schemas import RubricInput


def pdf():
    import importlib.util
    from pathlib import Path

    spec = importlib.util.spec_from_file_location("test_pdf", Path("scripts/create_test_pdf.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.paper_bytes()


@pytest.fixture
def assessment_case(owner, api_client, accounts):
    account, workspace, classroom = owner
    student = accounts()
    headers = auth(account)
    preview = api_client.post(
        f"/v1/classrooms/{classroom['id']}/imports/preview",
        headers=headers,
        files={
            "file": (
                "roster.csv",
                f"registration_number,name,email\nS001,Student,{student['email']}\nS002,Missing,missing@example.com",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    result = api_client.post(
        f"/v1/classrooms/{classroom['id']}/imports/{preview.json()['id']}/confirm", headers=headers
    )
    assert result.status_code == 200, result.text
    roster = api_client.get(f"/v1/classrooms/{classroom['id']}/students", headers=headers).json()
    invitation = api_client.post(
        f"/v1/classrooms/{classroom['id']}/invitations",
        headers=headers,
        json={"email": student["email"], "role": "student", "student_id": roster[0]["id"]},
    ).json()
    token = invitation["url"].split("#token=")[1]
    assert (
        api_client.post(
            "/v1/invitations/accept", headers=auth(student), json={"token": token}
        ).status_code
        == 200
    )
    result = api_client.post(
        f"/v1/classrooms/{classroom['id']}/assessments", headers=headers, json={"title": "CAT 1"}
    )
    assert result.status_code == 201, result.text
    assessment = result.json()
    result = api_client.post(
        f"/v1/assessments/{assessment['id']}/rubrics",
        headers=headers,
        json={
            "questions": [
                {
                    "label": "Q1",
                    "prompt": "Explain stack ordering.",
                    "maximum": "5.00",
                    "model_answer": "A stack uses LIFO.",
                    "criteria": [
                        {"label": "Identifies LIFO", "maximum": "3.00"},
                        {"label": "Explains last in, first out", "maximum": "2.00"},
                    ],
                }
            ]
        },
    )
    assert result.status_code == 201, result.text
    rubric = result.json()
    result = api_client.post(f"/v1/assessments/{assessment['id']}/activate", headers=headers)
    assert result.status_code == 200, result.text
    file = api_client.post(
        f"/v1/classrooms/{classroom['id']}/files",
        headers=headers,
        files={"file": ("answer.pdf", pdf(), "application/pdf")},
    ).json()
    result = api_client.post(
        f"/v1/assessments/{assessment['id']}/submissions",
        headers=headers,
        json={"student_id": roster[0]["id"], "file_id": file["id"], "identity_confirmed": True},
    )
    assert result.status_code == 201, result.text
    return account, classroom, assessment, rubric, result.json(), student


def review(client, account, submission, rubric, readable=True, base=0):
    return client.post(
        f"/v1/submissions/{submission['id']}/extractions",
        headers=auth(account),
        json={
            "base_version": base,
            "pages_complete": True,
            "identity_confirmed": True,
            "answers": [
                {
                    "question_id": rubric["questions"][0]["id"],
                    "text": "A stack uses LIFO: last in, first out.",
                    "page_from": 1,
                    "page_to": 1,
                    "readable": readable,
                }
            ],
        },
    )


def test_extraction_version_wins_over_clock_order(assessment_case, api_client, environment):
    account, classroom, assessment, rubric, paper, pupil = assessment_case
    first = review(api_client, account, paper, rubric)
    assert first.status_code == 201
    with psycopg.connect(
        environment.migration_database_url.replace("postgresql+psycopg://", "postgresql://")
    ) as db:
        db.execute(
            "UPDATE chalkmate.extractions SET created_at=now()+interval '1 day' WHERE id=%s",
            (first.json()["id"],),
        )
    second = review(api_client, account, paper, rubric, base=1)
    assert second.status_code == 201, second.text
    current = api_client.get(f"/v1/submissions/{paper['id']}", headers=auth(account))
    assert current.status_code == 200
    assert current.json()["extraction"]["id"] == second.json()["id"]
    assert current.json()["extraction"]["version"] == 2


def decision_body(extraction, rubric):
    return {
        "extraction_id": extraction["id"],
        "feedback": "Good understanding of LIFO.",
        "reason": "Manually checked against both rubric criteria.",
        "marks": [
            {"criterion_id": criterion["id"], "mark": "2.50" if i == 0 else "2.00"}
            for i, criterion in enumerate(rubric["questions"][0]["criteria"])
        ],
    }


def test_assessment_review_approve_release_and_student_privacy(
    assessment_case, api_client, environment, accounts
):
    account, classroom, assessment, rubric, paper, student = assessment_case
    assert (
        api_client.get(f"/v1/classrooms/{classroom['id']}/my-results", headers=auth(student)).json()
        == []
    )
    assert (
        api_client.get(f"/v1/submissions/{paper['id']}", headers=auth(student)).status_code == 404
    )
    extraction = review(api_client, account, paper, rubric)
    assert extraction.status_code == 201, extraction.text
    result = api_client.post(
        f"/v1/submissions/{paper['id']}/decisions",
        headers=auth(account),
        json=decision_body(extraction.json(), rubric),
    )
    assert result.status_code == 201, result.text
    assert result.json()["total"] == "4.50"
    assert (
        api_client.get(f"/v1/classrooms/{classroom['id']}/my-results", headers=auth(student)).json()
        == []
    )
    result = api_client.post(
        f"/v1/submissions/{paper['id']}/release",
        headers=auth(account),
        json={"decision_id": result.json()["id"]},
    )
    assert result.status_code == 200, result.text
    results = api_client.get(f"/v1/classrooms/{classroom['id']}/my-results", headers=auth(student))
    assert results.status_code == 200, results.text
    assert results.json()[0]["total"] == "4.50"
    assert results.json()[0]["questions"][0]["model_answer"] == ""
    assert review(api_client, account, paper, rubric, base=1).status_code == 409
    other = accounts()
    assert (
        api_client.get(f"/v1/assessments/{assessment['id']}", headers=auth(other)).status_code
        == 404
    )
    with psycopg.connect(
        environment.database_url.replace("postgresql+psycopg://", "postgresql://")
    ) as db:
        db.execute("SELECT set_config('app.user_id',%s,true)", (student["id"],))
        for table in ("extractions", "proposals", "decisions", "submissions", "rubrics"):
            assert db.execute(f"SELECT id FROM chalkmate.{table}").fetchall() == []
    exported = api_client.get(f"/v1/assessments/{assessment['id']}/export", headers=auth(account))
    assert "4.50" in exported.text and "missing" in exported.text
    excel = api_client.get(
        f"/v1/assessments/{assessment['id']}/export?format=xlsx", headers=auth(account)
    )
    assert excel.status_code == 200
    assert load_workbook(io.BytesIO(excel.content)).active.max_row == 3


def test_unreadable_and_stale_extractions_block_approval(assessment_case, api_client):
    account, _, _, rubric, paper, _ = assessment_case
    first = review(api_client, account, paper, rubric, readable=False)
    assert first.status_code == 201
    assert (
        api_client.post(
            f"/v1/submissions/{paper['id']}/decisions",
            headers=auth(account),
            json=decision_body(first.json(), rubric),
        ).status_code
        == 409
    )
    second = review(api_client, account, paper, rubric, base=1)
    assert second.status_code == 201
    assert review(api_client, account, paper, rubric, base=1).status_code == 409
    body = decision_body(second.json(), rubric)
    body["marks"][0]["mark"] = "3.01"
    assert (
        api_client.post(
            f"/v1/submissions/{paper['id']}/decisions", headers=auth(account), json=body
        ).status_code
        == 422
    )


def test_durable_pdf_job_is_idempotent(assessment_case, api_client):
    from chalkmate import processing

    account, _, _, _, paper, _ = assessment_case
    result = api_client.post(
        f"/v1/submissions/{paper['id']}/jobs",
        headers=auth(account),
        json={"kind": "extract", "model": ""},
    )
    assert result.status_code == 202, result.text
    job_id = result.json()["id"]
    processing.process_job.run(job_id)
    processing.process_job.run(job_id)
    result = api_client.get(f"/v1/submissions/{paper['id']}", headers=auth(account)).json()
    assert result["jobs"][0]["status"] == "succeeded"
    assert result["jobs"][0]["attempts"] == 1
    assert "LIFO" in result["extraction"]["pages"][0]["text"]


def test_correction_invalidates_approved_result(assessment_case, api_client):
    account, _, _, rubric, paper, _ = assessment_case
    extraction = review(api_client, account, paper, rubric).json()
    approved = api_client.post(
        f"/v1/submissions/{paper['id']}/decisions",
        headers=auth(account),
        json=decision_body(extraction, rubric),
    )
    assert approved.status_code == 201, approved.text
    assert review(api_client, account, paper, rubric, base=1).status_code == 201
    result = api_client.post(
        f"/v1/submissions/{paper['id']}/release",
        headers=auth(account),
        json={"decision_id": approved.json()["id"]},
    )
    assert result.status_code == 409


def test_override_requires_reason_and_preserves_original_proposal(assessment_case, api_client):
    from sqlalchemy.orm import Session

    from chalkmate.db import get_engine, set_identity

    account, classroom, _, rubric, paper, _ = assessment_case
    extraction = review(api_client, account, paper, rubric).json()
    criteria = rubric["questions"][0]["criteria"]
    with Session(get_engine()) as db, db.begin():
        set_identity(db, account["id"], account["email"])
        proposal = m.Proposal(
            workspace_id=UUID(classroom["workspace_id"]),
            classroom_id=UUID(classroom["id"]),
            submission_id=UUID(paper["id"]),
            extraction_id=UUID(extraction["id"]),
            rubric_id=UUID(rubric["id"]),
            model="synthetic-test",
            model_digest="synthetic-test",
            prompt_version="test",
            output={
                "criteria": [
                    {
                        "criterion_id": criterion["id"],
                        "question_id": rubric["questions"][0]["id"],
                        "mark": criterion["maximum"],
                        "explanation": "Test fixture",
                        "evidence_page": 1,
                        "evidence_quote": "LIFO",
                    }
                    for criterion in criteria
                ],
                "feedback": "Test fixture",
                "review_reasons": [],
            },
        )
        db.add(proposal)
        db.flush()
        proposal_id = str(proposal.id)
    body = decision_body(extraction, rubric)
    body["proposal_id"], body["reason"] = proposal_id, ""
    assert (
        api_client.post(
            f"/v1/submissions/{paper['id']}/decisions", headers=auth(account), json=body
        ).status_code
        == 422
    )
    body["reason"] = "Partial credit after checking the original paper."
    assert (
        api_client.post(
            f"/v1/submissions/{paper['id']}/decisions", headers=auth(account), json=body
        ).status_code
        == 201
    )
    latest = api_client.get(f"/v1/submissions/{paper['id']}", headers=auth(account)).json()
    assert latest["proposal"]["output"]["criteria"][0]["mark"] == "3.00"
    assert latest["decision"]["marks"][0]["mark"] == "2.50"


def test_rubric_score_and_evidence_validation():
    qid, cid = str(uuid4()), str(uuid4())
    questions = [{"id": qid, "criteria": [{"id": cid, "maximum": "5"}]}]
    answers = [{"question_id": qid, "text": "LIFO", "page_from": 1, "page_to": 1}]
    with pytest.raises(HTTPException):
        validate_marks(questions, [{"criterion_id": cid, "mark": "5.01"}])
    with pytest.raises(ValidationError):
        RubricInput.model_validate(
            {
                "questions": [
                    {
                        "label": "Q1",
                        "prompt": "Question",
                        "maximum": 5,
                        "criteria": [{"label": "Only criterion", "maximum": 4}],
                    }
                ]
            }
        )
    output = {
        "criteria": [
            {
                "criterion_id": cid,
                "question_id": qid,
                "mark": "4.5",
                "explanation": "Correct concept",
                "evidence_quote": "invented",
                "evidence_page": 1,
            }
        ],
        "feedback": "Feedback",
        "review_reasons": [],
    }
    with pytest.raises(ValueError):
        validate_suggestion(questions, answers, output)
    output["criteria"][0]["evidence_quote"] = "LIFO"
    assert validate_suggestion(questions, answers, output)["criteria"][0]["mark"] == "4.5"


@pytest.mark.skipif(os.environ.get("RUN_LOCAL_MODEL_TESTS") != "1", reason="Requires local Ollama")
def test_live_local_grading_job(assessment_case, api_client):
    from chalkmate.processing import process_job

    account, _, _, rubric, paper, _ = assessment_case
    extracted = review(api_client, account, paper, rubric)
    assert extracted.status_code == 201
    queued = api_client.post(
        f"/v1/submissions/{paper['id']}/jobs",
        headers=auth(account),
        json={"kind": "grade", "model": "gemma3:4b"},
    )
    assert queued.status_code == 202, queued.text
    process_job.run(queued.json()["id"])
    deadline = time.monotonic() + 120
    while time.monotonic() < deadline:
        result = api_client.get(f"/v1/submissions/{paper['id']}", headers=auth(account)).json()
        if result["jobs"][0]["status"] not in ("queued", "running"):
            break
        time.sleep(1)
    assert result["jobs"][0]["status"] == "succeeded", result["jobs"]
    assert result["proposal"]["model"] == "gemma3:4b"
    assert len(result["proposal"]["output"]["criteria"]) == 2
    assert result["status"] == "teacher_review"
