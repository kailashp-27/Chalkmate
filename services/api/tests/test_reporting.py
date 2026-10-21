"""Saved exports and teacher-reviewed bank mappings use real scoped PostgreSQL data."""

import csv
import io

import psycopg
import pytest
from conftest import auth
from openpyxl import load_workbook
from test_assessments import decision_body, review
from test_planning import ok, paper

pytest_plugins = ["test_assessments"]


@pytest.fixture
def reporting_case(assessment_case, api_client):
    account, classroom, exam, rubric, submission, student = assessment_case
    path = f"/v1/classrooms/{classroom['id']}/planning"
    source = ok(
        api_client.post(
            path,
            headers=auth(account),
            json={
                "kind": "syllabus",
                "title": "Stacks",
                "content": {
                    "topics": [
                        {
                            "name": "Stacks",
                            "text": "A stack uses LIFO ordering. Push and pop change the stack.",
                        }
                    ]
                },
            },
        ),
        201,
    )
    ok(api_client.post(f"/v1/planning/{source['id']}/approve", headers=auth(account)))
    bank = ok(
        api_client.post(
            path, headers=auth(account), json={**paper(source), "kind": "question_bank"}
        ),
        201,
    )
    ok(api_client.post(f"/v1/planning/{bank['id']}/approve", headers=auth(account)))
    template_body = {
        "kind": "marks_template",
        "title": "Department sheet",
        "content": {
            "score_basis": "approved",
            "export_columns": [
                {"field": "registration_number", "heading": "Roll no."},
                {"field": "status", "heading": "Review state"},
                {"field": "total", "heading": "Internal mark"},
                {"field": "name", "heading": "=Unsafe header"},
                {"field": "maximum", "heading": "Out of"},
                {"field": "percentage", "heading": "Percentage"},
            ],
        },
    }
    template = ok(api_client.post(path, headers=auth(account), json=template_body), 201)
    ok(api_client.post(f"/v1/planning/{template['id']}/approve", headers=auth(account)))
    return assessment_case, bank, template, template_body


def approve_marks(client, case, marks=None):
    account, classroom, exam, rubric, submission, student = case
    extraction = ok(review(client, account, submission, rubric), 201)
    body = decision_body(extraction, rubric)
    if marks:
        for row, mark in zip(body["marks"], marks, strict=True):
            row["mark"] = mark
    return ok(
        client.post(
            f"/v1/submissions/{submission['id']}/decisions", headers=auth(account), json=body
        ),
        201,
    ), extraction


def test_export_snapshot_columns_precision_privacy_and_immutability(
    reporting_case, api_client, environment, accounts
):
    case, bank, template, _ = reporting_case
    account, classroom, exam, rubric, submission, student = case
    decision, extraction = approve_marks(api_client, case)
    with psycopg.connect(
        environment.migration_database_url.replace("postgresql+psycopg://", "postgresql://")
    ) as db:
        db.execute(
            "UPDATE chalkmate.students SET name=%s WHERE id=%s",
            ("\t=1+1", submission["student_id"]),
        )
    path = f"/v1/assessments/{exam['id']}/marks-exports"
    snapshot = ok(
        api_client.post(path, headers=auth(account), json={"template_id": template["id"]}), 201
    )
    row = snapshot["snapshot"]["rows"][0]
    assert row["status"] == "approved" and row["total"] == "4.50" and row["percentage"] == "90.00"
    assert row["decision_id"] == decision["id"]
    output = api_client.get(f"/v1/marks-exports/{snapshot['id']}/download", headers=auth(account))
    assert output.status_code == 200
    rows = list(csv.reader(io.StringIO(output.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["Roll no.", "Review state", "Internal mark"]
    assert rows[0][3] == "'=Unsafe header" and rows[1][3] == "'=1+1"
    excel = api_client.get(
        f"/v1/marks-exports/{snapshot['id']}/download?format=xlsx", headers=auth(account)
    )
    sheet = load_workbook(io.BytesIO(excel.content)).active
    assert sheet.cell(2, 3).value == 4.5 and sheet.cell(2, 3).number_format == "0.00"
    assert sheet.cell(2, 4).data_type == "s"
    changed = decision_body(extraction, rubric)
    for mark in changed["marks"]:
        mark["mark"] = "1"
    ok(
        api_client.post(
            f"/v1/submissions/{submission['id']}/decisions", headers=auth(account), json=changed
        ),
        201,
    )
    assert (
        ok(api_client.get(f"/v1/marks-exports/{snapshot['id']}", headers=auth(account)))[
            "snapshot"
        ]["rows"][0]["total"]
        == "4.50"
    )
    fresh = ok(
        api_client.post(path, headers=auth(account), json={"template_id": template["id"]}), 201
    )
    assert fresh["snapshot"]["rows"][0]["total"] == "2.00"
    other = accounts()
    assert (
        api_client.get(
            f"/v1/marks-exports/{snapshot['id']}/download", headers=auth(other)
        ).status_code
        == 404
    )
    assert (
        api_client.get(f"/v1/marks-exports/{snapshot['id']}", headers=auth(student)).status_code
        == 404
    )
    with psycopg.connect(
        environment.database_url.replace("postgresql+psycopg://", "postgresql://")
    ) as db:
        db.execute("SELECT set_config('app.user_id',%s,true)", (other["id"],))
        assert db.execute("SELECT count(*) FROM chalkmate.marks_exports").fetchone()[0] == 0
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            db.execute(
                "UPDATE chalkmate.marks_exports SET digest='altered' WHERE id=%s", (snapshot["id"],)
            )


def test_export_missing_scores_and_template_revision(reporting_case, api_client):
    case, bank, template, body = reporting_case
    account, classroom, exam, rubric, submission, student = case
    path = f"/v1/assessments/{exam['id']}/marks-exports"
    preview = ok(
        api_client.post(path, headers=auth(account), json={"template_id": template["id"]}), 201
    )
    row = preview["snapshot"]["rows"][0]
    assert row["total"] is None and row["status"] == "unapproved"
    ok(
        api_client.post(
            f"/v1/planning/{template['id']}/versions",
            headers=auth(account),
            json={**body, "base_version": 1},
        ),
        201,
    )
    assert (
        api_client.post(
            path, headers=auth(account), json={"template_id": template["id"]}
        ).status_code
        == 409
    )
    assert (
        api_client.get(
            f"/v1/marks-exports/{preview['id']}/download", headers=auth(account)
        ).status_code
        == 200
    )


def test_question_links_measure_only_current_approved_answers(reporting_case, api_client):
    case, bank, _, _ = reporting_case
    account, classroom, exam, rubric, submission, student = case
    path = f"/v1/planning/{bank['id']}/question-links"
    body = {
        "question_index": 0,
        "assessment_id": exam["id"],
        "rubric_id": rubric["id"],
        "question_id": rubric["questions"][0]["id"],
        "reason": "Teacher reviewed this reuse.",
    }
    linked = ok(api_client.post(path, headers=auth(account), json=body), 201)
    assert linked["links"][0]["reviewed"] == 0 and linked["links"][0]["average"] is None
    assert api_client.post(path, headers=auth(account), json=body).status_code == 409
    approve_marks(api_client, case)
    current = ok(
        api_client.get(f"/v1/planning/{bank['id']}/question-performance", headers=auth(account))
    )
    link = current["links"][0]
    assert link["reviewed"] == 1 and link["average"] == "4.50" and not link["stale"]
    assert "student" not in str(current).lower()
    revoked = ok(
        api_client.post(
            path,
            headers=auth(account),
            json={
                **body,
                "base_link_id": link["id"],
                "disposition": "revoked",
                "reason": "Question changed.",
            },
        ),
        201,
    )
    assert not revoked["links"] and len(revoked["history"]) == 2
    assert all(r["reviewed"] == 0 for r in revoked["history"])


def test_stale_bank_and_cross_class_links_are_rejected(reporting_case, api_client, accounts):
    case, bank, _, _ = reporting_case
    account, classroom, exam, rubric, submission, student = case
    path = f"/v1/planning/{bank['id']}/question-links"
    body = {
        "question_index": 0,
        "assessment_id": exam["id"],
        "rubric_id": rubric["id"],
        "question_id": rubric["questions"][0]["id"],
        "reason": "Teacher linked the source.",
    }
    ok(api_client.post(path, headers=auth(account), json=body), 201)
    ok(
        api_client.post(
            f"/v1/planning/{bank['id']}/versions",
            headers=auth(account),
            json={
                "kind": "question_bank",
                "title": bank["title"],
                "content": bank["content"],
                "source_ids": [s["id"] for s in bank["sources"]],
                "base_version": 1,
            },
        ),
        201,
    )
    current = ok(
        api_client.get(f"/v1/planning/{bank['id']}/question-performance", headers=auth(account))
    )
    assert current["links"][0]["stale"] and current["links"][0]["average"] is None
    assert (
        api_client.post(
            path, headers=auth(account), json={**body, "base_link_id": current["links"][0]["id"]}
        ).status_code
        == 409
    )
    assert (
        api_client.get(
            f"/v1/planning/{bank['id']}/question-performance", headers=auth(student)
        ).status_code
        == 404
    )


def test_template_checks_missing_required_columns(reporting_case, api_client):
    case, _, _, body = reporting_case
    account, classroom, *_ = case
    body["content"]["export_columns"] = [
        {"field": "total", "heading": "Mark"},
        {"field": "total", "heading": "Mark"},
    ]
    created = ok(
        api_client.post(
            f"/v1/classrooms/{classroom['id']}/planning", headers=auth(account), json=body
        ),
        201,
    )
    assert len(created["issues"]) == 2
    assert (
        api_client.post(f"/v1/planning/{created['id']}/approve", headers=auth(account)).status_code
        == 409
    )
