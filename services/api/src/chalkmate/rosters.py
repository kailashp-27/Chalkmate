import csv
import io
import zipfile
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import EmailStr, TypeAdapter, ValidationError

MAX_ROSTER_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
email_adapter = TypeAdapter(EmailStr)


def parse_roster(filename: str, content: bytes) -> tuple[list[dict], list[dict]]:
    if len(content) > MAX_ROSTER_BYTES:
        raise HTTPException(413, "Roster must be 2 MB or smaller")
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".csv":
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            values = []
            for index, row in enumerate(reader):
                if index > MAX_ROWS:
                    raise HTTPException(422, "Import at most 1,000 students at a time")
                values.append(row)
        elif extension == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 16 * 1024 * 1024:
                    raise HTTPException(413, "Expanded spreadsheet exceeds the safety limit")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                sheet = workbook.active
                if sheet.max_column and sheet.max_column > 30:
                    raise HTTPException(422, "Use a roster with at most 30 columns")
                values = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index > MAX_ROWS:
                        raise HTTPException(422, "Import at most 1,000 students at a time")
                    values.append(list(row))
            finally:
                workbook.close()
        else:
            raise HTTPException(415, "Choose a CSV or XLSX roster")
    except (UnicodeError, csv.Error, zipfile.BadZipFile, ValueError, KeyError, OSError) as exc:
        raise HTTPException(
            422, "The roster could not be read. Use UTF-8 CSV or a valid XLSX."
        ) from exc
    if not values:
        raise HTTPException(422, "The roster is empty")
    headers = [str(value or "").strip().lower() for value in values[0]]
    required = ["registration_number", "name", "email"]
    if any(headers.count(key) != 1 for key in required):
        raise HTTPException(422, "Include each header once: registration_number, name, email")
    positions = {key: headers.index(key) for key in required}
    rows, issues, seen = [], [], set()
    for number, values_row in enumerate(values[1:], 2):
        if not any(value is not None and str(value).strip() for value in values_row):
            continue
        row = {
            key: str(
                values_row[index]
                if index < len(values_row) and values_row[index] is not None
                else ""
            ).strip()
            for key, index in positions.items()
        }
        row["registration_number"] = row["registration_number"].upper()
        row["email"] = row["email"].lower()

        def issue(message, severity="error"):
            issues.append({"row": number, "severity": severity, "message": message})

        if not row["registration_number"] or not row["name"]:
            issue("Registration number and name are required")
        if len(row["registration_number"]) > 80 or len(row["name"]) > 160:
            issue("Registration number or name is too long")
        if any(value.startswith(("=", "+", "-", "@")) for value in row.values()):
            issue("Formula-like values are not accepted; replace them with plain text")
        if row["registration_number"] in seen:
            issue("Duplicate registration number in this file")
        seen.add(row["registration_number"])
        try:
            row["email"] = str(email_adapter.validate_python(row["email"]))
        except ValidationError:
            issue(
                "Missing or invalid email; this student will be imported without email", "warning"
            )
            row["email"] = None
        rows.append({"row": number, **row})
    if not rows:
        raise HTTPException(422, "The roster has no student rows")
    return rows, issues
