"""Reviewable export snapshots and explicitly linked aggregate question performance."""

import csv
import io
from decimal import ROUND_HALF_UP, Decimal
from uuid import UUID

from fastapi import APIRouter, HTTPException, Response
from openpyxl import Workbook
from sqlalchemy import select, text

from chalkmate import assessment_models as grades
from chalkmate import planning_models as planning
from chalkmate import reporting_models as m
from chalkmate import reporting_schemas as s
from chalkmate.assessment_routes import assessment, rubric_for
from chalkmate.auth import CurrentUser, Database
from chalkmate.export_cells import safe_cell
from chalkmate.intelligence_routes import lock
from chalkmate.intelligence_rules import cohort, digest, summary
from chalkmate.planning_routes import approved, latest, output, record
from chalkmate.routes import audit, classroom

router = APIRouter(prefix="/v1")


def export_record(db, export_id):
    row = db.get(m.MarksExport, export_id)
    if not row:
        raise HTTPException(404, "Export not found")
    classroom(db, row.classroom_id)
    return row


@router.post(
    "/assessments/{assessment_id}/marks-exports", response_model=s.ExportOutput, status_code=201
)
def prepare_export(assessment_id: UUID, body: s.ExportRequest, db: Database, user: CurrentUser):
    exam = assessment(db, assessment_id, edit=True)
    template = record(db, body.template_id)
    if (
        template.classroom_id != exam.classroom_id
        or template.kind != "marks_template"
        or not approved(db, template)
        or not latest(db, template)
        or output(db, template)["issues"]
    ):
        raise HTTPException(409, "Choose a current approved marks template from this classroom")
    rubric, cohort_rows = cohort(db, exam)
    # A prior release stays available during re-evaluation, with an explicit status.
    releases = {}
    for paper, decision, extraction in db.execute(
        select(grades.Submission, grades.Decision, grades.Extraction)
        .join(grades.Release, grades.Release.submission_id == grades.Submission.id)
        .join(grades.Decision, grades.Decision.id == grades.Release.decision_id)
        .join(grades.Extraction, grades.Extraction.id == grades.Decision.extraction_id)
        .where(grades.Submission.assessment_id == exam.id)
        .order_by(grades.Extraction.version.desc(), grades.Release.created_at.desc())
    ):
        releases.setdefault(paper.student_id, (decision, extraction))
    rows = []
    content = template.content
    for entry in cohort_rows:
        withdrawn = entry["disposition"] == "withdrawn"
        if withdrawn and not content.get("include_withdrawn", False):
            continue
        decision = None
        if content.get("score_basis", "released") == "released":
            released = releases.get(entry["student"].id)
            if withdrawn:
                status = "withdrawn"
            elif released:
                decision, extracted = released
                status = (
                    "released"
                    if rubric
                    and decision.rubric_id == rubric.id
                    and entry["extraction"]
                    and extracted.id == entry["extraction"].id
                    else "released_previous"
                )
            else:
                status = "missing" if entry["disposition"] == "missing" else "not_released"
        else:
            status = entry["disposition"]
            if status == "reviewed":
                decision, status = entry["decision"], "approved"
        student = entry["student"]
        rows.append(
            s.ExportRow(
                registration_number=student.registration_number,
                name=student.name,
                enrollment="withdrawn" if withdrawn else "active",
                status=status,
                total=decision.total if decision else None,
                maximum=decision.maximum if decision else None,
                percentage=(decision.total / decision.maximum * 100).quantize(
                    Decimal(".01"), rounding=ROUND_HALF_UP
                )
                if decision and decision.maximum
                else None,
                feedback=decision.feedback if decision else "",
                decision_id=decision.id if decision else None,
                rubric_id=decision.rubric_id if decision else None,
            )
        )
    snapshot = s.ExportSnapshot(
        assessment_title=exam.title,
        template_title=template.title,
        template_version=template.version,
        score_basis=content.get("score_basis", "released"),
        include_withdrawn=content.get("include_withdrawn", False),
        columns=content["export_columns"],
        rows=rows,
    ).model_dump(mode="json")
    row = m.MarksExport(
        workspace_id=exam.workspace_id,
        classroom_id=exam.classroom_id,
        assessment_id=exam.id,
        template_id=template.id,
        created_by=user.id,
        snapshot=snapshot,
        digest=digest(snapshot),
    )
    db.add(row)
    db.flush()
    audit(
        db,
        user,
        exam.workspace_id,
        "marks_export.prepared",
        exam.classroom_id,
        export_id=str(row.id),
        digest=row.digest,
        template_id=str(template.id),
    )
    return row


@router.get("/classrooms/{class_id}/marks-exports", response_model=list[s.ExportOutput])
def list_exports(class_id: UUID, db: Database, user: CurrentUser):
    classroom(db, class_id)
    return list(
        db.scalars(
            select(m.MarksExport)
            .where(m.MarksExport.classroom_id == class_id)
            .order_by(m.MarksExport.created_at.desc())
            .limit(50)
        )
    )


@router.get("/marks-exports/{export_id}", response_model=s.ExportOutput)
def get_export(export_id: UUID, db: Database, user: CurrentUser):
    return export_record(db, export_id)


def export_bytes(snapshot, format):
    columns = snapshot["columns"]
    fields = [c["field"] for c in columns]
    rows = [[safe_cell(c["heading"]) for c in columns]]
    for row in snapshot["rows"]:
        rows.append([safe_cell(row[field]) if row[field] is not None else "" for field in fields])
    if format == "csv":
        stream = io.StringIO()
        csv.writer(stream).writerows(rows)
        return stream.getvalue().encode("utf-8-sig"), "text/csv"
    book = Workbook()
    sheet = book.active
    sheet.title = "Marks snapshot"
    sheet.append(rows[0])
    for row in rows[1:]:
        sheet.append(
            [
                Decimal(str(v))
                if fields[i] in ("total", "maximum", "percentage") and v != ""
                else v
                for i, v in enumerate(row)
            ]
        )
    for i, field in enumerate(fields, 1):
        if field in ("total", "maximum", "percentage"):
            for cells in sheet.iter_rows(min_row=2, min_col=i, max_col=i):
                cells[0].number_format = "0.00"
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = 24
    sheet.freeze_panes = "A2"
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/marks-exports/{export_id}/download")
def download(export_id: UUID, db: Database, user: CurrentUser, format: str = "csv"):
    row = export_record(db, export_id)
    if format not in ("csv", "xlsx"):
        raise HTTPException(422, "Choose CSV or XLSX")
    if digest(row.snapshot) != row.digest:
        raise HTTPException(409, "Export integrity check failed; prepare another preview")
    db.execute(
        text("SELECT chalkmate.record_marks_download(:id,:format)"),
        {"id": row.id, "format": format},
    )
    data, mime = export_bytes(row.snapshot, format)
    return Response(
        data,
        media_type=mime,
        headers={
            "Content-Disposition": f'attachment; filename="marks-{row.id}.{format}"',
            "Cache-Control": "no-store",
        },
    )


def link_key(row):
    return row.question_index, row.assessment_id, row.rubric_id, row.question_id


def bank(db, document_id, edit=False):
    doc = record(db, document_id, edit=edit)
    if doc.kind not in ("paper", "question_bank"):
        raise HTTPException(422, "Choose a question bank or paper")
    return doc


@router.post(
    "/planning/{document_id}/question-links", response_model=s.PerformanceOutput, status_code=201
)
def link_question(document_id: UUID, body: s.LinkInput, db: Database, user: CurrentUser):
    doc = bank(db, document_id, edit=True)
    lock(db, doc.id)
    if body.question_index >= len(doc.content.get("questions", [])):
        raise HTTPException(422, "Choose a question in this version")
    exam = assessment(db, body.assessment_id)
    rubric = db.get(grades.Rubric, body.rubric_id)
    if exam.classroom_id != doc.classroom_id or not rubric or rubric.assessment_id != exam.id:
        raise HTTPException(422, "Choose a rubric from this classroom")
    if not any(q["id"] == str(body.question_id) for q in rubric.questions):
        raise HTTPException(422, "Choose a question in the selected rubric")
    rows = list(
        db.scalars(
            select(m.QuestionLink)
            .where(m.QuestionLink.document_id == doc.id)
            .order_by(m.QuestionLink.version.desc())
        )
    )
    previous = next(
        (
            r
            for r in rows
            if link_key(r)
            == (body.question_index, body.assessment_id, body.rubric_id, body.question_id)
        ),
        None,
    )
    if body.base_link_id != (previous.id if previous else None):
        raise HTTPException(409, "Link history changed. Refresh before saving.")
    if body.disposition == "linked":
        if (
            not latest(db, doc)
            or not approved(db, doc)
            or output(db, doc)["issues"]
            or not rubric_for(db, exam)
            or rubric_for(db, exam).id != rubric.id
            or not rubric.approved_at
        ):
            raise HTTPException(409, "Link current approved bank and rubric versions only")
    elif not previous or previous.disposition != "linked":
        raise HTTPException(409, "There is no active link to revoke")
    db.add(
        m.QuestionLink(
            workspace_id=doc.workspace_id,
            classroom_id=doc.classroom_id,
            document_id=doc.id,
            version=max((r.version for r in rows), default=0) + 1,
            question_index=body.question_index,
            assessment_id=exam.id,
            rubric_id=rubric.id,
            question_id=body.question_id,
            disposition=body.disposition,
            reason=body.reason,
            created_by=user.id,
        )
    )
    audit(
        db,
        user,
        doc.workspace_id,
        "question.link_reviewed",
        doc.classroom_id,
        document_id=str(doc.id),
        disposition=body.disposition,
    )
    db.flush()
    return performance(document_id, db, user)


@router.get("/planning/{document_id}/question-performance", response_model=s.PerformanceOutput)
def performance(document_id: UUID, db: Database, user: CurrentUser):
    doc = bank(db, document_id)
    records = list(
        db.scalars(
            select(m.QuestionLink)
            .where(m.QuestionLink.document_id == doc.id)
            .order_by(m.QuestionLink.version.desc())
        )
    )
    current, summaries, history = {}, {}, []
    valid_bank = latest(db, doc) and approved(db, doc) and not output(db, doc)["issues"]
    for row in records:
        key = link_key(row)
        current.setdefault(key, row.id)
        exam = db.get(grades.Assessment, row.assessment_id)
        if exam.id not in summaries:
            summaries[exam.id] = summary(db, exam)
        data = summaries[exam.id]
        q = next((v for v in data["topics"] if v["question_id"] == str(row.question_id)), None)
        stale = not valid_bank or data["rubric_id"] != row.rubric_id or q is None
        measured = q if not stale and row.disposition == "linked" and current[key] == row.id else {}
        history.append(
            {
                "id": row.id,
                "question_index": row.question_index,
                "assessment_id": exam.id,
                "assessment_title": exam.title,
                "rubric_id": row.rubric_id,
                "question_id": row.question_id,
                "disposition": row.disposition,
                "reason": row.reason,
                "created_at": row.created_at,
                "stale": stale,
                "reviewed": measured.get("reviewed", 0),
                "maximum": measured.get("maximum"),
                "average": measured.get("average"),
                "full_credit": measured.get("full_credit", 0),
                "partial_credit": measured.get("partial_credit", 0),
                "zero_credit": measured.get("zero_credit", 0),
            }
        )
    reused = []
    for other in db.scalars(
        select(planning.PlanningDocument).where(
            planning.PlanningDocument.classroom_id == doc.classroom_id
        )
    ):
        if any(src.get("id") == str(doc.id) for src in other.sources):
            reused.append(
                {
                    "id": other.id,
                    "title": other.title,
                    "kind": other.kind,
                    "version": other.version,
                    "approved": approved(db, other),
                }
            )
    return {
        "links": [
            r for r in history if r["id"] in current.values() and r["disposition"] == "linked"
        ],
        "history": history,
        "reused_by": reused,
    }
